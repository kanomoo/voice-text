"""
Academic Project Report Template Engine (KMUTNB / Thai University Standard)
Full-featured production engine supporting two-pass TOC sync, zero dangling punctuation,
APA table formatting, native Word XML page numbering, PyMuPDF scanning, and cross-platform headless conversion.
"""

import os
import sys
import re
import subprocess
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from pythainlp.tokenize import word_tokenize

FONT_NAME = 'TH SarabunPSK'
CLR_BLACK = RGBColor(0, 0, 0)

# Characters that must never be separated from preceding text by a line-break or ZWSP
NO_BREAK_BEFORE = r'[,\.\:\;\!\?\)\]\}\%\"\'\sฯๆ/”’]'
# Characters that must never be followed by a line-break or ZWSP immediately after
NO_BREAK_AFTER = r'[\(\[\{\s\"\'/“‘]'

def thai_zwsp(text):
    """
    Inserts zero-width spaces (\u200b) between Thai syllables using PyThaiNLP
    while strictly removing ZWSP adjacent to punctuation, quotes, brackets, and percent signs
    to enforce the Zero Dangling Punctuation Rule.
    """
    if not text:
        return ""
    lines = str(text).split('\n')
    processed = []
    for line in lines:
        if not line:
            processed.append("")
            continue
        tokens = word_tokenize(line, engine='newmm')
        joined = '\u200b'.join(tokens)
        # Remove ZWSP before closing punctuation, brackets, quotes, %
        joined = re.sub(r'\u200b+(' + NO_BREAK_BEFORE + ')', r'\1', joined)
        # Remove ZWSP after opening punctuation, brackets, quotes
        joined = re.sub('(' + NO_BREAK_AFTER + r')\u200b+', r'\1', joined)
        # Clean any consecutive ZWSP
        joined = re.sub(r'\u200b+', '\u200b', joined)
        processed.append(joined)
    return '\n'.join(processed)

def set_run_font(run, size_pt=16, bold=False, italic=False, underline=False, color_rgb=CLR_BLACK):
    """Applies TH SarabunPSK font settings across all OpenXML script tags (ascii, hAnsi, cs, eastAsia)."""
    run.font.name = FONT_NAME
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    run.underline = underline
    run.font.color.rgb = color_rgb
    
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn('w:ascii'), FONT_NAME)
    rFonts.set(qn('w:hAnsi'), FONT_NAME)
    rFonts.set(qn('w:cs'), FONT_NAME)
    rFonts.set(qn('w:eastAsia'), FONT_NAME)
    
    sz_val = int(round(size_pt * 2))
    for sz in rPr.findall(qn('w:sz')):
        rPr.remove(sz)
    for szCs in rPr.findall(qn('w:szCs')):
        rPr.remove(szCs)
    rPr.append(parse_xml(f'<w:sz {nsdecls("w")} w:val="{sz_val}"/>'))
    rPr.append(parse_xml(f'<w:szCs {nsdecls("w")} w:val="{sz_val}"/>'))

    for b in rPr.findall(qn('w:b')):
        rPr.remove(b)
    for bCs in rPr.findall(qn('w:bCs')):
        rPr.remove(bCs)
    if bold:
        rPr.append(parse_xml(f'<w:b {nsdecls("w")}/>'))
        rPr.append(parse_xml(f'<w:bCs {nsdecls("w")}/>'))
    else:
        rPr.append(parse_xml(f'<w:b {nsdecls("w")} w:val="0"/>'))
        rPr.append(parse_xml(f'<w:bCs {nsdecls("w")} w:val="0"/>'))

    for i_elem in rPr.findall(qn('w:i')):
        rPr.remove(i_elem)
    for iCs in rPr.findall(qn('w:iCs')):
        rPr.remove(iCs)
    if italic:
        rPr.append(parse_xml(f'<w:i {nsdecls("w")}/>'))
        rPr.append(parse_xml(f'<w:iCs {nsdecls("w")}/>'))

def format_paragraph(p, space_before=0, space_after=2, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.THAI_JUSTIFY, keep_with_next=False):
    """Standard paragraph formatting helper with widow/orphan control and optional keep_with_next."""
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    p.alignment = align
    p.paragraph_format.widow_control = True
    if keep_with_next:
        p.paragraph_format.keep_with_next = True

def add_body_p(doc, text, bold_prefix=None, indent=True, space_after=3, page_break_before=False):
    """Standard body paragraph with 0.5-inch first-line indent and Thai Justify."""
    p = doc.add_paragraph()
    format_paragraph(p, space_before=0, space_after=space_after, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.THAI_JUSTIFY)
    p.paragraph_format.first_line_indent = Inches(0.5) if indent else Inches(0)
    p.paragraph_format.left_indent = Inches(0)
    if page_break_before:
        p.paragraph_format.page_break_before = True
    if bold_prefix:
        r_bold = p.add_run(thai_zwsp(bold_prefix))
        set_run_font(r_bold, size_pt=16, bold=True)
    r_text = p.add_run(thai_zwsp(text))
    set_run_font(r_text, size_pt=16, bold=False)
    return p

def add_numbered_item(doc, num_label, text, bold_prefix=None, space_after=2, keep_with_next=False, page_break_before=False):
    """
    Standard Thai academic numbered list (ตามระเบียบงานสารบรรณไทย).
    บรรทัดที่ 2 เป็นต้นไป เยื้อง 1 Tab (left_indent = 0.5 นิ้ว)
    บรรทัดที่ 1 ย่อหน้าเข้าไปมากกว่า (first_line_indent = 0.3 นิ้ว รวมเป็น 0.8 นิ้ว)
    """
    p = doc.add_paragraph()
    format_paragraph(p, space_before=0, space_after=space_after, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.THAI_JUSTIFY, keep_with_next=keep_with_next)
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(0.3)
    if page_break_before:
        p.paragraph_format.page_break_before = True
    
    r_num = p.add_run(f"{num_label}  ")
    set_run_font(r_num, size_pt=16, bold=False)
    if bold_prefix:
        r_bold = p.add_run(thai_zwsp(bold_prefix))
        set_run_font(r_bold, size_pt=16, bold=True)
    r_text = p.add_run(thai_zwsp(text))
    set_run_font(r_text, size_pt=16, bold=False)
    return p

def add_bullet_item(doc, text, bold_prefix=None, space_after=2, keep_with_next=False, page_break_before=False):
    """Standard Thai academic bullet item (left_indent = 0.5 in, first_line_indent = 0.3 in)."""
    p = doc.add_paragraph()
    format_paragraph(p, space_before=0, space_after=space_after, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.THAI_JUSTIFY, keep_with_next=keep_with_next)
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(0.3)
    if page_break_before:
        p.paragraph_format.page_break_before = True
    
    r_sym = p.add_run("•  ")
    set_run_font(r_sym, size_pt=16, bold=True)
    if bold_prefix:
        r_bold = p.add_run(thai_zwsp(bold_prefix))
        set_run_font(r_bold, size_pt=16, bold=True)
    r_text = p.add_run(thai_zwsp(text))
    set_run_font(r_text, size_pt=16, bold=False)
    return p

def add_subbullet_item(doc, text, bold_prefix=None, space_after=2, keep_with_next=False, page_break_before=False):
    """Standard Thai academic sub-bullet item (left_indent = 0.8 in, first_line_indent = 0.3 in)."""
    p = doc.add_paragraph()
    format_paragraph(p, space_before=0, space_after=space_after, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.THAI_JUSTIFY, keep_with_next=keep_with_next)
    p.paragraph_format.left_indent = Inches(0.8)
    p.paragraph_format.first_line_indent = Inches(0.3)
    if page_break_before:
        p.paragraph_format.page_break_before = True
    
    r_sym = p.add_run("-  ")
    set_run_font(r_sym, size_pt=16, bold=False)
    if bold_prefix:
        r_bold = p.add_run(thai_zwsp(bold_prefix))
        set_run_font(r_bold, size_pt=16, bold=True)
    r_text = p.add_run(thai_zwsp(text))
    set_run_font(r_text, size_pt=16, bold=False)
    return p

def add_heading_1(doc, text, page_break_before=False, space_before=12, space_after=3):
    """Heading 1: 16pt Bold, flush left, keep_with_next = True."""
    p = doc.add_paragraph()
    format_paragraph(p, space_before=space_before, space_after=space_after, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.LEFT, keep_with_next=True)
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.left_indent = Inches(0)
    if page_break_before:
        p.paragraph_format.page_break_before = True
    r = p.add_run(thai_zwsp(text))
    set_run_font(r, size_pt=16, bold=True)
    return p

def add_heading_2(doc, text, page_break_before=False, space_before=8, space_after=2, indent_inches=0.0):
    """Heading 2: 16pt Bold, keep_with_next = True."""
    p = doc.add_paragraph()
    format_paragraph(p, space_before=space_before, space_after=space_after, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.LEFT, keep_with_next=True)
    p.paragraph_format.first_line_indent = Inches(indent_inches)
    p.paragraph_format.left_indent = Inches(0)
    if page_break_before:
        p.paragraph_format.page_break_before = True
    r = p.add_run(thai_zwsp(text))
    set_run_font(r, size_pt=16, bold=True)
    return p

def add_heading_3(doc, text, page_break_before=False, space_before=6, space_after=2, indent_inches=0.5):
    """Heading 3: 16pt Bold, keep_with_next = True."""
    p = doc.add_paragraph()
    format_paragraph(p, space_before=space_before, space_after=space_after, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.LEFT, keep_with_next=True)
    p.paragraph_format.first_line_indent = Inches(indent_inches)
    p.paragraph_format.left_indent = Inches(0)
    if page_break_before:
        p.paragraph_format.page_break_before = True
    r = p.add_run(thai_zwsp(text))
    set_run_font(r, size_pt=16, bold=True)
    return p

def add_chapter_title(doc, chapter_num, title_text):
    """Official Chapter Heading: 'บทที่ X' in 20pt bold, followed by chapter title in 18pt bold, centered."""
    p = doc.add_paragraph()
    format_paragraph(p, space_before=14, space_after=14, align=WD_ALIGN_PARAGRAPH.CENTER, keep_with_next=True)
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.left_indent = Inches(0)
    r_num = p.add_run(f"บทที่ {chapter_num}\n")
    set_run_font(r_num, size_pt=20, bold=True)
    r_tit = p.add_run(thai_zwsp(title_text))
    set_run_font(r_tit, size_pt=18, bold=True)
    return p

def add_front_matter_title(doc, title_text, space_before=14, space_after=14, page_break_before=False):
    """Front matter or back matter major heading (e.g. 'บทคัดย่อ', 'สารบัญ', 'บรรณานุกรม') in 20pt bold, centered."""
    p = doc.add_paragraph()
    format_paragraph(p, space_before=space_before, space_after=space_after, align=WD_ALIGN_PARAGRAPH.CENTER, keep_with_next=True)
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.left_indent = Inches(0)
    if page_break_before:
        p.paragraph_format.page_break_before = True
    r = p.add_run(thai_zwsp(title_text))
    set_run_font(r, size_pt=20, bold=True)
    return p

def add_styled_table(doc, table_num, caption, headers, data, col_widths, font_size_pt=12.0, padding_twips=25, page_break_before=False):
    """
    Standard APA 3-Line Academic Table:
    - Caption above table, flush left, keep_with_next = True
    - Top line 1.5pt (sz=12), Bottom line 1.5pt (sz=12), Header underline 1.0pt (sz=8)
    - Zero vertical borders
    - cantSplit on all rows, tblHeader on header row
    - Cell margins/padding controlled in twips
    """
    p_cap = doc.add_paragraph()
    format_paragraph(p_cap, space_before=10, space_after=4, align=WD_ALIGN_PARAGRAPH.LEFT, keep_with_next=True)
    p_cap.paragraph_format.first_line_indent = Inches(0)
    if page_break_before:
        p_cap.paragraph_format.page_break_before = True
        
    r_lbl = p_cap.add_run(thai_zwsp(f"ตารางที่ {table_num}  "))
    set_run_font(r_lbl, size_pt=14, bold=True)
    r_cap = p_cap.add_run(thai_zwsp(caption))
    set_run_font(r_cap, size_pt=14, bold=False)

    tbl = doc.add_table(rows=len(data) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="12" w:space="0" w:color="000000"/>'
        f'  <w:bottom w:val="single" w:sz="12" w:space="0" w:color="000000"/>'
        f'  <w:insideH w:val="none"/>'
        f'  <w:insideV w:val="none"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tbl._tbl.tblPr.append(tblBorders)

    for i, row in enumerate(tbl.rows):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        if i == 0:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

    for c_idx, h_text in enumerate(headers):
        cell = tbl.cell(0, c_idx)
        cell.width = col_widths[c_idx]
        p = cell.paragraphs[0]
        format_paragraph(p, space_before=3, space_after=3, line_spacing=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
        p.paragraph_format.first_line_indent = Inches(0)
        r = p.add_run(thai_zwsp(h_text))
        set_run_font(r, size_pt=font_size_pt, bold=True)
        tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:bottom w:val="single" w:sz="8" w:space="0" w:color="000000"/></w:tcBorders>')
        cell._tc.get_or_add_tcPr().append(tcBorders)

    for r_idx, row_values in enumerate(data):
        for c_idx, val in enumerate(row_values):
            cell = tbl.cell(r_idx + 1, c_idx)
            cell.width = col_widths[c_idx]
            p = cell.paragraphs[0]
            format_paragraph(p, space_before=2, space_after=2, line_spacing=1.0, align=WD_ALIGN_PARAGRAPH.LEFT)
            p.paragraph_format.first_line_indent = Inches(0)
            r = p.add_run(thai_zwsp(str(val)))
            set_run_font(r, size_pt=font_size_pt, bold=False)

    for row in tbl.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tcMar = parse_xml(
                f'<w:tcMar {nsdecls("w")}>'
                f'  <w:top w:w="{padding_twips}" w:type="dxa"/>'
                f'  <w:bottom w:w="{padding_twips}" w:type="dxa"/>'
                f'  <w:left w:w="40" w:type="dxa"/>'
                f'  <w:right w:w="40" w:type="dxa"/>'
                f'</w:tcMar>'
            )
            tcPr.append(tcMar)
    return tbl

def add_figure(doc, img_path, fig_num, caption, width_inches=3.0, source_text=None, page_break_before=False):
    """Adds figure centered with proper academic caption underneath and optional citation."""
    if not os.path.exists(img_path):
        return None
    p_img = doc.add_paragraph()
    format_paragraph(p_img, space_before=10, space_after=4, align=WD_ALIGN_PARAGRAPH.CENTER, keep_with_next=True)
    p_img.paragraph_format.first_line_indent = Inches(0)
    if page_break_before:
        p_img.paragraph_format.page_break_before = True
    p_img.add_run().add_picture(img_path, width=Inches(width_inches))

    p_cap = doc.add_paragraph()
    format_paragraph(p_cap, space_before=2, space_after=4 if source_text else 10, align=WD_ALIGN_PARAGRAPH.CENTER, keep_with_next=bool(source_text))
    p_cap.paragraph_format.first_line_indent = Inches(0)
    r_lbl = p_cap.add_run(thai_zwsp(f"รูปที่ {fig_num}  "))
    set_run_font(r_lbl, size_pt=14, bold=True)
    r_cap = p_cap.add_run(thai_zwsp(caption))
    set_run_font(r_cap, size_pt=14, bold=False)

    if source_text:
        p_src = doc.add_paragraph()
        format_paragraph(p_src, space_before=0, space_after=10, align=WD_ALIGN_PARAGRAPH.CENTER)
        p_src.paragraph_format.first_line_indent = Inches(0)
        r_src = p_src.add_run(thai_zwsp(f"(ที่มา: {source_text})"))
        set_run_font(r_src, size_pt=12, italic=True)
    return p_cap

def add_equation(doc, eq_text, eq_num_str, space_before=6, space_after=6):
    """
    Mathematical Equation centered with right tab stop at 8300 dxa (5.76 in)
    for numbering '(X.X)'.
    """
    p = doc.add_paragraph()
    format_paragraph(p, space_before=space_before, space_after=space_after, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.CENTER)
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.left_indent = Inches(0)
    pPr = p._p.get_or_add_pPr()
    tab_xml = parse_xml(f'<w:tabs {nsdecls("w")}><w:tab w:val="right" w:pos="8300"/></w:tabs>')
    pPr.append(tab_xml)
    r_eq = p.add_run(eq_text)
    set_run_font(r_eq, size_pt=14, italic=True)
    r_num = p.add_run(f"\t({eq_num_str})")
    set_run_font(r_num, size_pt=14, bold=True)
    return p

def add_likert_summary_table(doc, table_num, caption, criteria_data, col_widths=None, overall_row=None):
    """
    Specialized helper for Likert Evaluation Results:
    Headers: ['ลำดับ', 'ประเด็นการประเมิน', 'X̄', 'S.D.', 'ระดับคุณภาพ']
    criteria_data is list of tuples: (item_title, mean_float, sd_float, interpretation_str)
    """
    headers = ["ลำดับ", "ประเด็นการประเมิน", "ค่าเฉลี่ย (X̄)", "ส่วนเบี่ยงเบน (S.D.)", "ระดับคุณภาพ"]
    if col_widths is None:
        col_widths = [Inches(0.6), Inches(2.8), Inches(0.9), Inches(1.0), Inches(1.1)]
    table_data = []
    for idx, (crit_name, mean_val, sd_val, level_str) in enumerate(criteria_data, 1):
        m_str = f"{mean_val:.2f}" if isinstance(mean_val, (int, float)) else str(mean_val)
        sd_str = f"{sd_val:.2f}" if isinstance(sd_val, (int, float)) else str(sd_val)
        table_data.append([str(idx), crit_name, m_str, sd_str, level_str])
    if overall_row:
        table_data.append(["", overall_row[0], overall_row[1], overall_row[2], overall_row[3]])
    return add_styled_table(doc, table_num, caption, headers, table_data, col_widths)

def add_reference_item(doc, ref_text):
    """
    Standard APA Hanging Indent Reference Item:
    First line: flush left (0 inches)
    Subsequent lines: indented 0.5 inches
    """
    p = doc.add_paragraph()
    format_paragraph(p, space_before=2, space_after=4, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.THAI_JUSTIFY)
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.5)
    r = p.add_run(thai_zwsp(ref_text))
    set_run_font(r, size_pt=15, bold=False)
    return p

def add_toc_line(doc, title_text, page_num_str, indent=0.0, is_bold=False):
    """
    Table of Contents line with Right Tab Stop at 8300 dxa (5.76 inches)
    and Dot Leaders (w:leader="dot") for 100% alignment.
    """
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Inches(indent)
    p.paragraph_format.first_line_indent = Inches(0)
    pPr = p._p.get_or_add_pPr()
    tab_xml = parse_xml(f'<w:tabs {nsdecls("w")}><w:tab w:val="right" w:leader="dot" w:pos="8300"/></w:tabs>')
    pPr.append(tab_xml)
    r1 = p.add_run(thai_zwsp(title_text))
    set_run_font(r1, size_pt=14, bold=is_bold)
    r2 = p.add_run(f"\t{page_num_str}")
    set_run_font(r2, size_pt=14, bold=is_bold)
    return p

def add_toc_col_header(doc, left_label="เรื่อง", right_label="หน้า"):
    """Adds the standard TOC column header row ('เรื่อง' - 'หน้า') aligned with right tab pos 8300 dxa."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Inches(0)
    p.paragraph_format.first_line_indent = Inches(0)
    pPr = p._p.get_or_add_pPr()
    tab_xml = parse_xml(f'<w:tabs {nsdecls("w")}><w:tab w:val="right" w:pos="8300"/></w:tabs>')
    pPr.append(tab_xml)
    r1 = p.add_run(thai_zwsp(left_label))
    set_run_font(r1, size_pt=14, bold=True)
    r2 = p.add_run(f"\t{right_label}")
    set_run_font(r2, size_pt=14, bold=True)
    return p

def add_header_page_field(header_obj):
    """Injects native Microsoft Word XML PAGE field into the top-right header."""
    h_p = header_obj.paragraphs[0]
    format_paragraph(h_p, space_before=0, space_after=0, align=WD_ALIGN_PARAGRAPH.RIGHT)
    h_p.paragraph_format.first_line_indent = Inches(0)
    h_p.paragraph_format.left_indent = Inches(0)
    h_run = h_p.add_run()
    set_run_font(h_run, size_pt=14)
    h_run._r.append(parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w')))
    h_run._r.append(parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w')))
    h_run._r.append(parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w')))
    h_run._r.append(parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w')))

def clear_pg_num_types(sec):
    """Removes pgNumType from section properties to ensure correct continuous page numbering."""
    for pnt in sec._sectPr.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pgNumType'):
        sec._sectPr.remove(pnt)

def setup_section_margins(sec, top=1.5, bottom=1.0, left=1.5, right=1.0, is_cover=False):
    """
    Standardizes section margin settings across documents in inches.
    When is_cover=True, left and right margins are set to 1.0 inch (Zero Margin Bias).
    """
    sec.top_margin = Inches(top)
    sec.bottom_margin = Inches(bottom)
    sec.left_margin = Inches(1.0 if is_cover else left)
    sec.right_margin = Inches(1.0 if is_cover else right)
    sec.header_distance = Inches(0.75)
    sec.footer_distance = Inches(0.5)
    if is_cover:
        sec.different_first_page_header_footer = True
    else:
        sec.header.is_linked_to_previous = False
        sec.footer.is_linked_to_previous = False

def build_cover_page(doc, logo_path, title_en, title_th, submission_lines, group_title, students):
    """
    Builds the official 1-page Cover with Zero Margin Bias (Left 1.0", Right 1.0", Top 1.5", Bottom 1.0")
    and an Invisible 3-Column Table for student names and IDs so columns align vertically 100%.
    """
    sec = doc.sections[0]
    sec.top_margin = Inches(1.5)
    sec.bottom_margin = Inches(1.0)
    sec.left_margin = Inches(1.0)
    sec.right_margin = Inches(1.0)
    sec.different_first_page_header_footer = True
    clear_pg_num_types(sec)

    if logo_path and os.path.exists(logo_path):
        p_logo = doc.add_paragraph()
        format_paragraph(p_logo, space_before=0, space_after=12, align=WD_ALIGN_PARAGRAPH.CENTER)
        p_logo.paragraph_format.first_line_indent = Inches(0)
        p_logo.add_run().add_picture(logo_path, width=Inches(1.25))

    p_title = doc.add_paragraph()
    format_paragraph(p_title, space_before=0, space_after=6, align=WD_ALIGN_PARAGRAPH.CENTER)
    p_title.paragraph_format.first_line_indent = Inches(0)
    if title_en:
        r_en = p_title.add_run(f"{title_en}\n")
        set_run_font(r_en, size_pt=15, bold=True)
    r_th = p_title.add_run(thai_zwsp(title_th))
    set_run_font(r_th, size_pt=18, bold=True)

    if submission_lines:
        p_sub = doc.add_paragraph()
        format_paragraph(p_sub, space_before=14, space_after=12, align=WD_ALIGN_PARAGRAPH.CENTER)
        p_sub.paragraph_format.first_line_indent = Inches(0)
        for i, line in enumerate(submission_lines):
            r = p_sub.add_run(thai_zwsp(line) + ("\n" if i < len(submission_lines)-1 else ""))
            set_run_font(r, size_pt=15, bold=False)

    p_grp = doc.add_paragraph()
    format_paragraph(p_grp, space_before=8, space_after=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    p_grp.paragraph_format.first_line_indent = Inches(0)
    r_g1 = p_grp.add_run("จัดทำโดย\n")
    set_run_font(r_g1, size_pt=16, bold=True)
    if group_title:
        r_g2 = p_grp.add_run(thai_zwsp(group_title))
        set_run_font(r_g2, size_pt=16, bold=True)

    if students:
        tbl = doc.add_table(rows=len(students), cols=3)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        tbl_borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="none"/><w:bottom w:val="none"/><w:left w:val="none"/>'
            f'<w:right w:val="none"/><w:insideH w:val="none"/><w:insideV w:val="none"/>'
            f'</w:tblBorders>'
        )
        tbl._tbl.tblPr.append(tbl_borders)
        col_widths = [Inches(1.5), Inches(1.7), Inches(1.5)]
        for r_idx, (pfx_first, last, sid) in enumerate(students):
            row = tbl.rows[r_idx]
            for c_idx, text in enumerate([pfx_first, last, sid]):
                cell = row.cells[c_idx]
                cell.width = col_widths[c_idx]
                p = cell.paragraphs[0]
                align = WD_ALIGN_PARAGRAPH.LEFT if c_idx < 2 else WD_ALIGN_PARAGRAPH.RIGHT
                format_paragraph(p, space_before=0, space_after=1, align=align)
                p.paragraph_format.first_line_indent = Inches(0)
                r = p.add_run(thai_zwsp(text))
                set_run_font(r, size_pt=14, bold=False)

def convert_to_pdf(doc_path, pdf_path):
    """
    Cross-platform conversion engine:
    Uses Microsoft Word / PowerPoint COM API on Windows,
    and headless LibreOffice ('soffice --headless --convert-to pdf') on Linux.
    """
    doc_path = os.path.abspath(doc_path)
    pdf_path = os.path.abspath(pdf_path)
    out_dir = os.path.dirname(pdf_path)

    if sys.platform == 'win32':
        import win32com.client
        if doc_path.lower().endswith('.docx'):
            word = win32com.client.Dispatch("Word.Application")
            doc = word.Documents.Open(doc_path)
            doc.SaveAs(pdf_path, FileFormat=17) # 17 = wdExportFormatPDF
            doc.Close(False)
            word.Quit()
        elif doc_path.lower().endswith('.pptx'):
            ppt = win32com.client.Dispatch("PowerPoint.Application")
            prs = ppt.Presentations.Open(doc_path, WithWindow=False)
            prs.SaveAs(pdf_path, 32) # 32 = ppSaveAsPDF
            prs.Close()
            ppt.Quit()
    else:
        cmd = ['soffice', '--headless', '--convert-to', 'pdf', doc_path, '--outdir', out_dir]
        subprocess.run(cmd, check=True)

def scan_pdf_toc_pages(pdf_path, search_patterns, body_start_page=7):
    """
    Scans a compiled PDF using PyMuPDF (fitz) to extract actual page numbers for TOC/LOT/LOF.
    `search_patterns` is a dict of {key: regex_pattern}.
    Returns a dict of {key: str(actual_page_number)}.
    """
    import fitz
    doc = fitz.open(pdf_path)
    results = {}
    for p_idx, page in enumerate(doc):
        text = page.get_text()
        for key, pattern in search_patterns.items():
            if key not in results and re.search(pattern, text):
                # Page number relative to main body start
                actual_page = (p_idx + 1) - body_start_page + 1
                results[key] = str(max(1, actual_page))
    doc.close()
    return results

def render_pdf_to_png(pdf_path, output_dir, pages=None, zoom=2.0):
    """
    Renders selected PDF pages to high-resolution PNG for automated visual inspection.
    `pages` is an optional list of 1-based page numbers.
    """
    import fitz
    os.makedirs(output_dir, exist_ok=True)
    doc = fitz.open(pdf_path)
    rendered_paths = []
    mat = fitz.Matrix(zoom, zoom)
    for p_idx, page in enumerate(doc):
        p_num = p_idx + 1
        if pages and p_num not in pages:
            continue
        pix = page.get_pixmap(matrix=mat)
        out_path = os.path.join(output_dir, f"page_{p_num:02d}.png")
        pix.save(out_path)
        rendered_paths.append(out_path)
    doc.close()
    return rendered_paths
